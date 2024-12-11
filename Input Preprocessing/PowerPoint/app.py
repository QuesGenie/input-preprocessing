from pptx import Presentation
from transformers import pipeline
from sentence_transformers import SentenceTransformer, util
import tqdm.notebook as tq
from tqdm.notebook import tqdm
import json
import uuid
from utils.extract import *
from candidate_labels import * 
from utils import utils



def PowerPointExtractor(path,relevance_model):
    folder_name = path.split('/')[-1].split('.')[0] + str(uuid.uuid4())
    folder_path = os.path.join('Data', folder_name)
    folder_image_path = os.path.join(folder_path, 'Images')
    folder_text_path = os.path.join(folder_path, 'Text')

    os.makedirs(folder_image_path, exist_ok=True)
    os.makedirs(folder_text_path, exist_ok=True)

    extracted_content = {"slides": []}
    ppt = Presentation(path)

    # Extract all text from slides for topic modeling
    all_text = extract_all_text(ppt)

    labels = generate_candidate_labels(all_text)
    print(labels)
    slide_number = 1

    for slide in tq.tqdm(ppt.slides):
        slide_data = {"slide_number": slide_number, "content": []}
        index = 1

        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                text = shape.text_frame.text.strip()
                preprocessed_text = preprocess_text(text)
                if preprocessed_text and is_relevant_text(preprocessed_text, relevance_model, labels):
                    slide_data["content"].append({
                        "type": "text",
                        "text": text,
                    })

            elif hasattr(shape, "image"):
                image_data = shape.image.blob
                # Extract text from the image
                ocr_text = extract_text_from_image(image_data)
                if not ocr_text:
                  image_path = os.path.join(folder_image_path, f"image_{slide_number}_{index}.png")
                  with open(image_path, 'wb') as f:
                      f.write(image_data)

                  slide_data["content"].append({
                      "type": "image",
                      "placeholder": f"[Image {slide_number}_{index}]",
                      "image_path": image_path
                  })
                  index += 1
                else:
                  slide_data["content"].append({
                      "type": "image",
                      "ocr_text": ocr_text 
                  })
                  index += 1


        if slide_data["content"]:
            extracted_content["slides"].append(slide_data)
        slide_number += 1

    file_name = os.path.join(folder_text_path, f'Data-{folder_name}.json')
    utils.write_json_to_file(extracted_content, file_name)
    return file_name


if __name__ == "__main__":
    relevance_model = pipeline("zero-shot-classification", model="valhalla/distilbart-mnli-12-1")
    file_name = PowerPointExtractor('/content/Lecture 8.pptx',relevance_model)
    print(f"Extraction complete. JSON saved at {file_name}")